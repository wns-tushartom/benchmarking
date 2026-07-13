from __future__ import annotations

from io import BytesIO
import hashlib
import json
from pathlib import Path
import shutil

from openpyxl import Workbook
import pytest

from source.services import project_questions
from source.services.project_questions import (
    ProjectQuestionStorageError,
    ProjectQuestionValidationError,
    create_question_set,
    load_normalized_questions,
    parse_question_bytes,
    parse_typed_question,
)
from source.services.project_workspace import (
    ProjectWorkspace,
    UploadLimits,
    validate_resource_id,
)


def _xlsx_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def _sparse_xlsx_bytes(last_row: int) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.cell(row=1, column=1, value="question")
    sheet.cell(row=last_row, column=1, value="Refund policy?")
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def test_typed_question_is_canonical_and_always_evidence_only():
    questions = parse_typed_question("How do I request a refund?")

    assert [question.to_row() for question in questions] == [
        {
            "schema_version": 1,
            "question_id": "q_000001",
            "query": "How do I request a refund?",
            "labels": {
                "reference_contexts": [],
                "source_ids": [],
                "chunk_refs": [],
            },
            "reference_answer": None,
            "source_row": 1,
        }
    ]


def test_answer_only_csv_remains_evidence_only():
    questions = parse_question_bytes(
        "questions.csv",
        b"question,answer\nrefund policy,ask support\n",
    )

    assert questions[0].reference_answer == "ask support"
    assert questions[0].labels.is_empty()
    assert questions[0].source_row == 2


def test_csv_normalizes_retrieval_labels_and_requires_qualified_chunk_refs():
    source_id = "source_" + "a" * 64
    questions = parse_question_bytes(
        "questions.csv",
        (
            "question,reference_context,source_id,chunk_ref,ground_truth\n"
            f'What is covered?,"refund policy; cancellation fee",{source_id},'
            'fixed_tok1200_ov150:42;fixed_tok1200_ov150:42,Reference answer\n'
        ).encode("utf-8"),
    )

    row = questions[0].to_row()
    assert row["labels"] == {
        "reference_contexts": ["refund policy", "cancellation fee"],
        "source_ids": [source_id],
        "chunk_refs": [
            {"chunker": "fixed_tok1200_ov150", "chunk_id": "42"},
        ],
    }
    assert row["reference_answer"] == "Reference answer"
    assert row["source_row"] == 2

    with pytest.raises(ProjectQuestionValidationError, match="qualified"):
        parse_question_bytes("questions.csv", b"question,chunk_ref\nrefund,42\n")


def test_txt_and_xlsx_question_modes_preserve_physical_source_rows():
    txt = parse_question_bytes("questions.txt", b"First question\n\nSecond question\n")
    assert [(question.query, question.source_row) for question in txt] == [
        ("First question", 1),
        ("Second question", 3),
    ]
    assert all(question.labels.is_empty() for question in txt)

    xlsx = parse_question_bytes(
        "questions.xlsx",
        _xlsx_bytes(
            [
                ["query", "context", "answer"],
                ["Refund?", "refunds are accepted for thirty days", "Contact support"],
            ]
        ),
    )
    assert xlsx[0].query == "Refund?"
    assert xlsx[0].source_row == 2
    assert xlsx[0].labels.reference_contexts == (
        "refunds are accepted for thirty days",
    )
    assert xlsx[0].reference_answer == "Contact support"


def test_question_parsers_reject_duplicates_missing_columns_and_limits():
    with pytest.raises(ProjectQuestionValidationError, match="duplicate"):
        parse_question_bytes("questions.txt", b"Same question\nSame question\n")
    with pytest.raises(ProjectQuestionValidationError, match="question"):
        parse_question_bytes("questions.csv", b"answer\nNothing\n")
    with pytest.raises(ProjectQuestionValidationError, match="malformed"):
        parse_question_bytes("questions.csv", b'question\n"unterminated\n')
    with pytest.raises(ProjectQuestionValidationError, match="signature"):
        parse_question_bytes("questions.txt", b"%PDF-1.7\nnot really text questions\n")
    with pytest.raises(ProjectQuestionValidationError, match="signature"):
        parse_question_bytes("questions.csv", b"PK\x03\x04not really csv")
    with pytest.raises(ProjectQuestionValidationError, match="signature"):
        parse_question_bytes(
            "questions.txt",
            b"\x01\x02\x03binary-control-payload\n",
        )
    with pytest.raises(ProjectQuestionValidationError, match="4,000"):
        parse_typed_question("x" * 4001)
    with pytest.raises(ProjectQuestionValidationError, match="maximum"):
        parse_question_bytes("questions.txt", b"One\nTwo\n", max_questions=1)


def test_csv_and_xlsx_bound_physical_rows_before_skipping_blanks():
    sparse_csv = b"question\n" + (b"\n" * 10_000) + b"Refund policy?\n"
    with pytest.raises(ProjectQuestionValidationError, match="maximum"):
        parse_question_bytes("questions.csv", sparse_csv, max_questions=1)
    with pytest.raises(ProjectQuestionValidationError, match="maximum"):
        parse_question_bytes(
            "questions.xlsx",
            _sparse_xlsx_bytes(10_002),
            max_questions=1,
        )


def test_workspace_question_limits_are_loaded_from_environment(
    tmp_path: Path, monkeypatch
):
    workspace = ProjectWorkspace(tmp_path / "projects")
    project = workspace.create_upload("alpha.txt", b"Alpha project text", "Alpha")
    monkeypatch.setenv("PROJECT_MAX_QUESTIONS", "1")
    monkeypatch.setenv("PROJECT_MAX_QUESTION_FILE_KB", "1")

    with pytest.raises(ProjectQuestionValidationError, match="maximum question count"):
        workspace.create_question_set(
            project["project_id"],
            "questions.txt",
            b"Question one\nQuestion two\n",
        )
    with pytest.raises(ProjectQuestionValidationError, match="maximum question count"):
        workspace.create_question_set(
            project["project_id"],
            "questions.txt",
            b"Question one\nQuestion two\n",
            max_questions=2,
        )
    with pytest.raises(ProjectQuestionValidationError, match="size limit"):
        workspace.create_question_set(
            project["project_id"],
            "questions.txt",
            b"x" * 1025,
        )


def test_workspace_question_xlsx_uses_upload_archive_limits(tmp_path: Path):
    workspace = ProjectWorkspace(
        tmp_path / "projects",
        limits=UploadLimits(max_zip_entries=1),
    )
    project = workspace.create_upload("alpha.txt", b"Alpha project text", "Alpha")

    with pytest.raises(ProjectQuestionValidationError, match="container"):
        workspace.create_question_set(
            project["project_id"],
            "questions.xlsx",
            _xlsx_bytes([["question"], ["Refund policy?"]]),
        )


def test_invalid_workspace_question_limits_fail_closed(tmp_path: Path, monkeypatch):
    workspace = ProjectWorkspace(tmp_path / "projects")
    project = workspace.create_upload("alpha.txt", b"Alpha project text", "Alpha")
    monkeypatch.setenv("PROJECT_MAX_QUESTIONS", "not-a-number")

    with pytest.raises(ProjectQuestionValidationError, match="limit configuration"):
        workspace.create_question_set(
            project["project_id"],
            "questions.txt",
            b"Question one\n",
        )


def test_question_set_is_immutable_project_scoped_and_hash_verified(tmp_path: Path):
    workspace = ProjectWorkspace(tmp_path / "projects")
    alpha = workspace.create_upload("alpha.txt", b"Alpha project text", "Alpha")
    beta = workspace.create_upload("beta.txt", b"Beta project text", "Beta")

    created = workspace.create_question_set(
        alpha["project_id"],
        "questions.csv",
        b"question,answer\nRefund?,Ask support\n",
    )

    assert created["question_set_id"].startswith("questions_")
    assert validate_resource_id(created["question_set_id"]) == created["question_set_id"]
    question_root = (
        tmp_path
        / "projects"
        / alpha["project_id"]
        / "questions"
        / created["question_set_id"]
    )
    assert {path.name for path in question_root.iterdir()} == {
        "questions.csv",
        "questions.jsonl",
        "manifest.json",
    }
    loaded = workspace.load_question_set(
        alpha["project_id"],
        created["question_set_id"],
        expected_content_sha256=created["content_sha256"],
    )
    assert loaded[0].query == "Refund?"

    with pytest.raises((FileNotFoundError, ValueError)):
        workspace.load_question_set(
            beta["project_id"],
            created["question_set_id"],
            expected_content_sha256=created["content_sha256"],
        )

    raw = question_root / "questions.csv"
    raw.write_bytes(b"question\nTampered\n")
    with pytest.raises(ProjectQuestionValidationError, match="hash"):
        workspace.load_question_set(
            alpha["project_id"],
            created["question_set_id"],
            expected_content_sha256=created["content_sha256"],
        )


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("question_count", True),
        ("schema_version", 1.0),
    ],
)
def test_question_set_loader_rejects_noncanonical_manifest_types(
    tmp_path: Path,
    field: str,
    value: object,
):
    workspace = ProjectWorkspace(tmp_path / "projects")
    project = workspace.create_upload("alpha.txt", b"Alpha project text", "Alpha")
    created = workspace.create_question_set(
        project["project_id"],
        "questions.txt",
        b"Original question?\n",
    )
    manifest_path = (
        tmp_path
        / "projects"
        / project["project_id"]
        / "questions"
        / created["question_set_id"]
        / "manifest.json"
    )
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest[field] = value
    manifest_path.write_text(
        json.dumps(manifest, sort_keys=True, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )

    with pytest.raises(ProjectQuestionValidationError, match="manifest schema"):
        workspace.load_question_set(
            project["project_id"],
            created["question_set_id"],
            expected_content_sha256=created["content_sha256"],
        )


@pytest.mark.parametrize("mutation", ["schema_float", "nonsequential_id", "noncanonical_labels"])
def test_normalized_question_loader_requires_producer_canonical_schema(
    tmp_path: Path,
    mutation: str,
):
    row = parse_typed_question("Refund policy?")[0].to_row()
    if mutation == "schema_float":
        row["schema_version"] = 1.0
    elif mutation == "nonsequential_id":
        row["question_id"] = "q_000002"
    else:
        row["labels"]["reference_contexts"] = [" refund policy ", "refund policy"]
    target = tmp_path / "questions.jsonl"
    target.write_text(
        json.dumps(
            row,
            sort_keys=True,
            ensure_ascii=False,
            separators=(",", ":"),
        )
        + "\n",
        encoding="utf-8",
    )

    with pytest.raises(ProjectQuestionValidationError, match="canonical|schema"):
        load_normalized_questions(target)


def test_question_set_loader_rejects_normalized_and_manifest_tampering(tmp_path: Path):
    workspace = ProjectWorkspace(tmp_path / "projects")
    project = workspace.create_upload("alpha.txt", b"Alpha project text", "Alpha")
    created = workspace.create_question_set(
        project["project_id"],
        "questions.txt",
        b"Original question?\n",
    )
    root = (
        tmp_path
        / "projects"
        / project["project_id"]
        / "questions"
        / created["question_set_id"]
    )
    normalized_path = root / "questions.jsonl"
    normalized_row = json.loads(normalized_path.read_text(encoding="utf-8"))
    normalized_row["query"] = "Substituted question?"
    normalized = (
        json.dumps(
            normalized_row,
            sort_keys=True,
            ensure_ascii=False,
            separators=(",", ":"),
        ).encode("utf-8")
        + b"\n"
    )
    normalized_path.write_bytes(normalized)
    manifest_path = root / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["normalized_sha256"] = hashlib.sha256(normalized).hexdigest()
    manifest_path.write_text(
        json.dumps(manifest, sort_keys=True, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )

    with pytest.raises(ProjectQuestionValidationError, match="normalized"):
        workspace.load_question_set(
            project["project_id"],
            created["question_set_id"],
            expected_content_sha256=created["content_sha256"],
        )


def test_question_set_loader_rejects_directory_swap_to_other_project(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    workspace = ProjectWorkspace(tmp_path / "projects")
    alpha = workspace.create_upload("alpha.txt", b"Alpha project text", "Alpha")
    beta = workspace.create_upload("beta.txt", b"Beta project text", "Beta")
    created = workspace.create_question_set(
        alpha["project_id"],
        "questions.txt",
        b"Original question?\n",
    )
    alpha_root = (
        tmp_path
        / "projects"
        / alpha["project_id"]
        / "questions"
        / created["question_set_id"]
    )
    beta_root = (
        tmp_path
        / "projects"
        / beta["project_id"]
        / "questions"
        / created["question_set_id"]
    )
    shutil.copytree(alpha_root, beta_root)
    normalized_path = beta_root / "questions.jsonl"
    normalized_row = json.loads(normalized_path.read_text(encoding="utf-8"))
    normalized_row["query"] = "PROJECT B SECRET"
    normalized = (
        json.dumps(
            normalized_row,
            sort_keys=True,
            ensure_ascii=False,
            separators=(",", ":"),
        ).encode("utf-8")
        + b"\n"
    )
    normalized_path.write_bytes(normalized)
    manifest_path = beta_root / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["normalized_sha256"] = hashlib.sha256(normalized).hexdigest()
    manifest_path.write_text(
        json.dumps(manifest, sort_keys=True, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    original_reader = project_questions._read_regular_bytes_at
    swapped = False

    def swap_then_read(parent_fd: int, name: str, description: str) -> bytes:
        nonlocal swapped
        if not swapped and description == "question set manifest":
            swapped = True
            shutil.rmtree(alpha_root)
            alpha_root.symlink_to(beta_root, target_is_directory=True)
        return original_reader(parent_fd, name, description)

    monkeypatch.setattr(project_questions, "_read_regular_bytes_at", swap_then_read)
    with pytest.raises(ProjectQuestionStorageError):
        workspace.load_question_set(
            alpha["project_id"],
            created["question_set_id"],
            expected_content_sha256=created["content_sha256"],
        )


def test_question_set_publication_rolls_back_after_parent_fsync_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    questions_root = tmp_path / "questions"
    questions_root.mkdir()
    question_set_id = "questions_123e4567e89b42d3a456426614174000"
    root_identity = (questions_root.stat().st_dev, questions_root.stat().st_ino)
    real_fsync = project_questions.os.fsync
    failed_once = False

    def fail_parent_fsync(descriptor: int) -> None:
        nonlocal failed_once
        opened = project_questions.os.fstat(descriptor)
        if not failed_once and (opened.st_dev, opened.st_ino) == root_identity:
            failed_once = True
            raise OSError("forced parent fsync failure")
        real_fsync(descriptor)

    monkeypatch.setattr(project_questions.os, "fsync", fail_parent_fsync)
    with pytest.raises(ProjectQuestionStorageError):
        create_question_set(
            questions_root,
            question_set_id,
            "questions.txt",
            b"Refund policy?\n",
        )

    assert not (questions_root / question_set_id).exists()
    assert not list(questions_root.glob(".staging-*"))


def test_question_set_publication_reports_indeterminate_when_rollback_fsync_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    questions_root = tmp_path / "questions"
    questions_root.mkdir()
    question_set_id = "questions_123e4567e89b42d3a456426614174000"
    root_identity = (questions_root.stat().st_dev, questions_root.stat().st_ino)
    real_fsync = project_questions.os.fsync

    def fail_every_parent_fsync(descriptor: int) -> None:
        opened = project_questions.os.fstat(descriptor)
        if (opened.st_dev, opened.st_ino) == root_identity:
            raise OSError("forced parent fsync failure")
        real_fsync(descriptor)

    monkeypatch.setattr(project_questions.os, "fsync", fail_every_parent_fsync)
    with pytest.raises(ProjectQuestionStorageError, match="indeterminate"):
        create_question_set(
            questions_root,
            question_set_id,
            "questions.txt",
            b"Refund policy?\n",
        )

    assert not (questions_root / question_set_id).exists()
    assert not list(questions_root.glob(".staging-*"))


def test_direct_question_storage_rejects_non_resource_id(tmp_path: Path):
    questions_root = tmp_path / "questions"
    questions_root.mkdir()

    with pytest.raises(ProjectQuestionValidationError, match="question_set_id"):
        create_question_set(
            questions_root,
            "../escape",
            "questions.txt",
            b"Refund policy?\n",
        )
    assert not (tmp_path / "escape").exists()


def test_question_set_loader_rejects_symlinked_manifest(tmp_path: Path):
    workspace = ProjectWorkspace(tmp_path / "projects")
    project = workspace.create_upload("alpha.txt", b"Alpha project text", "Alpha")
    created = workspace.create_question_set(
        project["project_id"],
        "questions.txt",
        b"Refund policy?\n",
    )
    root = (
        tmp_path
        / "projects"
        / project["project_id"]
        / "questions"
        / created["question_set_id"]
    )
    real_manifest = root / "real-manifest.json"
    (root / "manifest.json").rename(real_manifest)
    (root / "manifest.json").symlink_to(real_manifest)

    with pytest.raises(ProjectQuestionStorageError, match="manifest"):
        workspace.load_question_set(
            project["project_id"],
            created["question_set_id"],
            expected_content_sha256=created["content_sha256"],
        )


def test_normalized_question_loader_rejects_symlinked_parent(tmp_path: Path):
    questions_root = tmp_path / "questions"
    questions_root.mkdir()
    created = create_question_set(
        questions_root,
        "questions_123e4567e89b42d3a456426614174000",
        "questions.txt",
        b"Refund policy?\n",
    )
    real = questions_root / created["question_set_id"]
    alias = tmp_path / "alias"
    alias.symlink_to(real, target_is_directory=True)

    with pytest.raises(ProjectQuestionStorageError):
        load_normalized_questions(alias / "questions.jsonl")


def test_normalized_question_loader_rejects_tampered_schema(tmp_path: Path):
    target = tmp_path / "questions.jsonl"
    target.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "question_id": "q_000001",
                "query": "Refund?",
                "labels": {
                    "reference_contexts": [],
                    "source_ids": [],
                    "chunk_refs": [],
                },
                "reference_answer": None,
                "source_row": 1,
                "unexpected": True,
            }
        )
        + "\n",
        encoding="utf-8",
    )
    with pytest.raises(ProjectQuestionValidationError, match="schema"):
        load_normalized_questions(target)
