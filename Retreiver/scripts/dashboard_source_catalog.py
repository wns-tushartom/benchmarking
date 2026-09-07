"""Allowlisted dataset and ground-truth discovery for the benchmark dashboard.

``build_source_catalog`` returns JSON-ready metadata only. Filesystem paths remain
private and are available solely through the resolver functions after a stable ID
has been matched against a freshly discovered allowlist.
"""
from __future__ import annotations

import csv
import json
import os
import shutil
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

DEFAULT_DATASET_ID = "dataset:wns-default"
NONE_GROUNDTRUTH_ID = "groundtruth:none"
CANONICAL_WORKBOOK = Path("data/chunking_methods_output_v2.xlsx")
BENCHMARK_CONFIG = Path("configs/benchmark.local.json")
TABLE_SUFFIXES = {".csv", ".xlsx"}

QUERY_COLUMNS = ("query", "question", "user_query", "prompt")
TEXT_COLUMNS = (
    "paragraph",
    "expected_text",
    "context",
    "ground_truth",
    "ground truth",
    "answer",
    "relevant_text",
    "groundtruth",
    "expected_answer",
    "reference_answer",
)
SOURCE_COLUMNS = (
    "pdf_name",
    "expected_pdf",
    "relevant_pdf",
    "document",
    "file",
    "filename",
    "source",
    "source_file",
    "source_doc",
)


@dataclass(frozen=True)
class DatasetSource:
    id: str
    label: str
    kind: str
    document_count: int
    chunk_count: int
    ready: bool
    validation: str
    sheets: tuple[str, ...] = ()
    chunk_counts_by_strategy: tuple[tuple[str, int], ...] = ()
    document_path: Path | None = None
    workbook_path: Path | None = None
    manifest_path: Path | None = None
    search_index_path: Path | None = None

    def public(self) -> dict[str, Any]:
        return {
            "id": self.id,
            "label": self.label,
            "kind": self.kind,
            "document_count": self.document_count,
            "chunk_count": self.chunk_count,
            "chunk_counts_by_strategy": dict(self.chunk_counts_by_strategy),
            "ready": self.ready,
            "validation": self.validation,
            "sheets": list(self.sheets),
        }


@dataclass(frozen=True)
class GroundtruthSource:
    id: str
    label: str
    kind: str
    row_count: int
    valid: bool
    validation: str
    query_column: str | None = None
    answer_column: str | None = None
    source_column: str | None = None
    project_id: str | None = None
    path: Path | None = None

    def public(self) -> dict[str, Any]:
        return {
            "id": self.id,
            "label": self.label,
            "kind": self.kind,
            "row_count": self.row_count,
            "valid": self.valid,
            "validation": self.validation,
            "query_column": self.query_column,
            "answer_column": self.answer_column,
            "source_column": self.source_column,
            "project_id": self.project_id,
        }


def _read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError):
        return {}
    return value if isinstance(value, dict) else {}


def _csv_metadata(path: Path) -> tuple[int, list[str]] | None:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            headers = next(reader, [])
            count = sum(1 for row in reader if any(str(cell).strip() for cell in row))
    except (OSError, UnicodeError, csv.Error):
        return None
    return count, [str(header).strip() for header in headers]


def _xlsx_sheets(path: Path) -> list[tuple[int, list[str]]] | None:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        return None
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets: list[tuple[int, list[str]]] = []
        try:
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows(values_only=True)
                header = next(rows, ())
                headers = ["" if value is None else str(value).strip() for value in header]
                count = sum(1 for row in rows if any(value is not None and str(value).strip() for value in row))
                sheets.append((count, headers))
        finally:
            workbook.close()
        return sheets
    except Exception:
        return None


def _workbook_sheet_names(path: Path | None) -> tuple[str, ...]:
    return tuple(_runnable_workbook_sheets(path)) if path is not None else ()


def _runnable_workbook_sheets(path: Path | None) -> list[str]:
    if path is None:
        return []
    try:
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            runnable: list[str] = []
            required = {"id", "pdf_name", "paragraph"}
            for worksheet in workbook.worksheets:
                first = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
                headers = {str(value).strip().lower() for value in first if value is not None}
                if required.issubset(headers):
                    runnable.append(str(worksheet.title))
            return runnable
        finally:
            workbook.close()
    except Exception:
        return []


def _workbook_chunk_counts(path: Path | None) -> tuple[tuple[str, int], ...]:
    if path is None:
        return ()
    try:
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            required = {"id", "pdf_name", "paragraph"}
            counts: list[tuple[str, int]] = []
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows(values_only=True)
                headers = {str(value).strip().lower() for value in next(rows, ()) if value is not None}
                if not required.issubset(headers):
                    continue
                count = sum(1 for row in rows if any(value is not None and str(value).strip() for value in row))
                counts.append((str(worksheet.title), count))
            return tuple(counts)
        finally:
            workbook.close()
    except Exception:
        return ()


def _configured_default_chunkers(root: Path) -> tuple[str, ...] | None:
    config_path = _safe_file(root, root / BENCHMARK_CONFIG)
    payload = _read_json(config_path) if config_path is not None else {}
    matrix = payload.get("matrix")
    raw_chunkers = matrix.get("chunkers") if isinstance(matrix, dict) else None
    if not isinstance(raw_chunkers, list):
        return None
    chunkers = tuple(dict.fromkeys(
        value.strip() for value in raw_chunkers
        if isinstance(value, str) and value.strip()
    ))
    return chunkers or None


def _table_metadata(path: Path) -> tuple[int, list[str]] | None:
    if path.suffix.lower() == ".csv":
        return _csv_metadata(path)
    if path.suffix.lower() == ".xlsx":
        sheets = _xlsx_sheets(path)
        return sheets[0] if sheets else None
    return None


def _default_chunk_count(root: Path) -> int:
    workbook = _safe_file(root, root / CANONICAL_WORKBOOK)
    sheets = _xlsx_sheets(workbook) if workbook is not None else None
    if sheets:
        return max((count for count, _headers in sheets), default=0)

    benchmark_input = _safe_file(root, root / "data" / "benchmark_input.csv")
    metadata = _csv_metadata(benchmark_input) if benchmark_input is not None else None
    if metadata is not None:
        return metadata[0]

    for candidate in (
        root / "data" / "chunks.json",
        root / "data" / "chunking_output.json",
        root / "data" / "search_index.json",
    ):
        safe_candidate = _safe_file(root, candidate)
        payload = _read_json(safe_candidate) if safe_candidate is not None else {}
        chunks = payload.get("chunks")
        if isinstance(chunks, list):
            return len(chunks)
        count = payload.get("chunk_count")
        if isinstance(count, int) and not isinstance(count, bool) and count >= 0:
            return count
    return 0


def _string_list(value: Any) -> list[str]:
    return [item for item in value if isinstance(item, str) and item] if isinstance(value, list) else []


def _inside(base: Path, candidate: Path) -> bool:
    try:
        candidate.resolve().relative_to(base.resolve())
        return True
    except (OSError, ValueError):
        return False


def _safe_directory(base: Path, candidate: Path) -> Path | None:
    try:
        lexical = Path(os.path.abspath(candidate))
        resolved = candidate.resolve(strict=True)
        approved_base = base.resolve(strict=True)
    except OSError:
        return None
    return resolved if resolved == lexical and resolved.is_dir() and _inside(approved_base, resolved) else None


def _safe_file(base: Path, candidate: Path) -> Path | None:
    try:
        lexical = Path(os.path.abspath(candidate))
        resolved = candidate.resolve(strict=True)
        approved_base = base.resolve(strict=True)
    except OSError:
        return None
    return resolved if resolved == lexical and resolved.is_file() and _inside(approved_base, resolved) else None


def _safe_declared_file(root: Path, base: Path, raw: Any) -> Path | None:
    if not isinstance(raw, str) or not raw.strip():
        return None
    value = Path(raw)
    candidates = [value] if value.is_absolute() else [root / value, base / value]
    for candidate in candidates:
        resolved = _safe_file(base, candidate)
        if resolved is not None:
            return resolved
    return None


def _project_document_count(root: Path, project_dir: Path, manifest: dict[str, Any], index: dict[str, Any]) -> int:
    declared_groups = (
        index.get("source_files"),
        manifest.get("source_files"),
        manifest.get("extracted_files"),
    )
    for value in declared_groups:
        files = _string_list(value)
        if files:
            resolved = {_safe_declared_file(root, project_dir, raw) for raw in files}
            return len({path for path in resolved if path is not None})
    saved = _safe_declared_file(root, project_dir, manifest.get("saved_path"))
    return 1 if saved is not None else 0


def _project_chunk_count(index: dict[str, Any]) -> int:
    chunks = index.get("chunks")
    if isinstance(chunks, list):
        return len(chunks)
    count = index.get("chunk_count")
    return count if isinstance(count, int) and not isinstance(count, bool) and count >= 0 else 0


def _project_workbook(root: Path, project_dir: Path, manifest: dict[str, Any], index: dict[str, Any]) -> Path | None:
    for raw in (index.get("workbook"), manifest.get("workbook"), "chunks/chunking_methods_output_v2.xlsx"):
        path = _safe_declared_file(root, project_dir, raw)
        if path is not None and path.suffix.lower() == ".xlsx" and _runnable_workbook_sheets(path):
            return path
    return None


def _detect_column(headers: list[str], candidates: tuple[str, ...]) -> str | None:
    normalized = {header.strip().lower(): header for header in headers if header.strip()}
    return next((normalized[name] for name in candidates if name in normalized), None)


def _table_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]] | None:
    if path.suffix.lower() == ".csv":
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                headers = [str(value).strip() for value in (reader.fieldnames or [])]
                return headers, [dict(row) for row in reader]
        except (OSError, UnicodeError, csv.Error):
            return None
    if path.suffix.lower() != ".xlsx":
        return None
    try:
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            if worksheet is None:
                return None
            values = worksheet.iter_rows(values_only=True)
            headers = ["" if value is None else str(value).strip() for value in next(values, ())]
            return headers, [dict(zip(headers, row)) for row in values]
        finally:
            workbook.close()
    except Exception:
        return None


def _groundtruth_record(
    source_id: str,
    label: str,
    kind: str,
    path: Path,
    *,
    project_id: str | None = None,
) -> GroundtruthSource:
    table = _table_rows(path)
    if table is None:
        return GroundtruthSource(source_id, label, kind, 0, False, "unreadable", project_id=project_id, path=path)
    headers, rows = table
    query = _detect_column(headers, QUERY_COLUMNS)
    answer = _detect_column(headers, TEXT_COLUMNS)
    source = _detect_column(headers, SOURCE_COLUMNS)
    query_rows = [row for row in rows if query and str(row.get(query) or "").strip()]
    row_count = len(query_rows)
    relevance_rows = [
        row for row in query_rows
        if (answer and str(row.get(answer) or "").strip()) or (source and str(row.get(source) or "").strip())
    ]
    valid = bool(row_count > 0 and query and (answer or source) and len(relevance_rows) == row_count)
    if not query:
        validation = "missing_query_column"
    elif not (answer or source):
        validation = "missing_relevance_column"
    elif row_count <= 0:
        validation = "empty"
    elif len(relevance_rows) != row_count:
        validation = "missing_relevance_values"
    else:
        validation = "valid"
    return GroundtruthSource(
        source_id,
        label,
        kind,
        row_count,
        valid,
        validation,
        query,
        answer,
        source,
        project_id,
        path,
    )


def _uploaded_groundtruth_path(root: Path, upload_dir: Path, manifest: dict[str, Any]) -> Path | None:
    for key in ("filename", "file", "path", "saved_path"):
        path = _safe_declared_file(root, upload_dir, manifest.get(key))
        if path is not None and path.suffix.lower() in TABLE_SUFFIXES:
            return path
    files = sorted(
        resolved
        for candidate in upload_dir.iterdir()
        if candidate.suffix.lower() in TABLE_SUFFIXES
        for resolved in [_safe_file(upload_dir, candidate)]
        if resolved is not None
    )
    return files[0] if len(files) == 1 else None


def _project_groundtruth_path(root: Path, project_dir: Path, manifest: dict[str, Any]) -> Path | None:
    path = _safe_declared_file(root, project_dir, manifest.get("question_file"))
    if path is not None and path.suffix.lower() in TABLE_SUFFIXES:
        return path
    questions_dir = _safe_directory(project_dir, project_dir / "questions")
    if questions_dir is None:
        return None
    files = sorted(
        resolved
        for candidate in questions_dir.iterdir()
        if candidate.suffix.lower() in TABLE_SUFFIXES
        for resolved in [_safe_file(questions_dir, candidate)]
        if resolved is not None
    )
    return files[0] if len(files) == 1 else None


def _discover_sources(root: Path) -> tuple[list[DatasetSource], list[GroundtruthSource]]:
    root = root.resolve()
    pdf_dir = _safe_directory(root, root / "data" / "pdfs")
    pdf_files = [] if pdf_dir is None else [
        resolved
        for candidate in pdf_dir.iterdir()
        if candidate.suffix.lower() == ".pdf"
        for resolved in [_safe_file(pdf_dir, candidate)]
        if resolved is not None
    ]
    document_count = len(pdf_files)
    chunk_count = _default_chunk_count(root)
    workbook_candidate = _safe_file(root, root / CANONICAL_WORKBOOK)
    workbook = workbook_candidate if workbook_candidate is not None and _runnable_workbook_sheets(workbook_candidate) else None
    workbook_sheets = _workbook_sheet_names(workbook)
    workbook_counts = _workbook_chunk_counts(workbook)
    configured_chunkers = _configured_default_chunkers(root)
    configured_sheets_present = True
    if configured_chunkers is not None:
        count_by_sheet = dict(workbook_counts)
        workbook_sheets = tuple(name for name in configured_chunkers if name in count_by_sheet)
        workbook_counts = tuple((name, count_by_sheet[name]) for name in workbook_sheets)
        configured_sheets_present = len(workbook_sheets) == len(configured_chunkers)
    default_ready = bool(
        document_count > 0
        and chunk_count > 0
        and workbook is not None
        and configured_sheets_present
    )
    datasets = [
        DatasetSource(
            DEFAULT_DATASET_ID,
            "WNS default dataset",
            "default",
            document_count,
            chunk_count,
            default_ready,
            "ready" if default_ready else "documents_chunks_or_workbook_missing",
            sheets=workbook_sheets,
            chunk_counts_by_strategy=workbook_counts,
            document_path=pdf_dir,
            workbook_path=workbook,
        )
    ]

    project_manifests: list[tuple[Path, str, dict[str, Any]]] = []
    projects_dir = _safe_directory(root, root / "data" / "user_projects")
    if projects_dir is not None:
        for manifest_candidate in sorted(projects_dir.glob("*/manifest.json")):
            original_project_dir = manifest_candidate.parent
            if original_project_dir.is_symlink():
                continue
            project_dir = _safe_directory(projects_dir, original_project_dir)
            if project_dir is None:
                continue
            manifest_path = _safe_file(project_dir, manifest_candidate)
            if manifest_path is None:
                continue
            manifest = _read_json(manifest_path)
            project_id = original_project_dir.name
            project_manifests.append((project_dir, project_id, manifest))
            index_path = _safe_file(project_dir, project_dir / "search_index.json")
            index = _read_json(index_path) if index_path is not None else {}
            document_count = _project_document_count(root, project_dir, manifest, index)
            chunk_count = _project_chunk_count(index)
            workbook_path = _project_workbook(root, project_dir, manifest, index)
            canonical_corpus = _safe_declared_file(
                root, project_dir, manifest.get("canonical_corpus")
            )
            raw_uploads = _safe_directory(project_dir, project_dir / "raw_uploads")
            document_path = raw_uploads or project_dir
            legacy_ready = bool(index and document_count > 0 and chunk_count > 0 and workbook_path)
            project_matrix_ready = bool(
                manifest.get("extraction_status") == "complete"
                and canonical_corpus is not None
                and canonical_corpus.suffix.lower() == ".jsonl"
                and document_count > 0
            )
            ready = legacy_ready or project_matrix_ready
            if legacy_ready:
                validation = "ready"
            elif project_matrix_ready:
                validation = "ready_for_project_matrix"
            else:
                validation = "manifest_corpus_or_documents_missing"
            datasets.append(
                DatasetSource(
                    f"project:{project_id}",
                    str(manifest.get("label") or project_id),
                    "uploaded_project",
                    document_count,
                    chunk_count,
                    ready,
                    validation,
                    sheets=_workbook_sheet_names(workbook_path),
                    chunk_counts_by_strategy=_workbook_chunk_counts(workbook_path),
                    document_path=document_path,
                    workbook_path=workbook_path,
                    manifest_path=manifest_path,
                    search_index_path=index_path,
                )
            )
    groundtruth: list[GroundtruthSource] = [
        GroundtruthSource(NONE_GROUNDTRUTH_ID, "None — evidence-only", "none", 0, True, "valid")
    ]
    repository_dir = _safe_directory(root, root / "data" / "groundtruth")
    if repository_dir is not None:
        for candidate in sorted(repository_dir.iterdir()):
            if candidate.suffix.lower() not in TABLE_SUFFIXES:
                continue
            path = _safe_file(repository_dir, candidate)
            if path is None:
                continue
            groundtruth.append(
                _groundtruth_record(
                    f"groundtruth:repository:{candidate.name}",
                    candidate.stem.replace("_", " "),
                    "repository",
                    path,
                )
            )

    uploads_dir = _safe_directory(root, root / "data" / "user_groundtruth")
    if uploads_dir is not None:
        for manifest_candidate in sorted(uploads_dir.glob("*/manifest.json")):
            original_upload_dir = manifest_candidate.parent
            if original_upload_dir.is_symlink():
                continue
            upload_dir = _safe_directory(uploads_dir, original_upload_dir)
            if upload_dir is None:
                continue
            manifest_path = _safe_file(upload_dir, manifest_candidate)
            if manifest_path is None:
                continue
            manifest = _read_json(manifest_path)
            path = _uploaded_groundtruth_path(root, upload_dir, manifest)
            source_id = f"groundtruth:upload:{original_upload_dir.name}"
            label = str(manifest.get("label") or original_upload_dir.name)
            groundtruth.append(
                _groundtruth_record(source_id, label, "uploaded", path)
                if path is not None
                else GroundtruthSource(source_id, label, "uploaded", 0, False, "file_missing_or_unsafe")
            )

    for project_dir, project_id, manifest in project_manifests:
        path = _project_groundtruth_path(root, project_dir, manifest)
        if path is None:
            continue
        label = str(manifest.get("label") or project_id)
        groundtruth.append(
            _groundtruth_record(
                f"groundtruth:project:{project_id}",
                f"{label} — uploaded ground truth",
                "project",
                path,
                project_id=project_id,
            )
        )

    return datasets, groundtruth


def register_groundtruth_upload(
    root: str | Path,
    original_name: str,
    content: bytes,
    label: str,
) -> dict[str, Any]:
    root_path = Path(root).resolve()
    filename = Path(original_name).name
    suffix = Path(filename).suffix.lower()
    if not filename or suffix not in TABLE_SUFFIXES:
        raise ValueError("Ground-truth upload must be a CSV or XLSX file")
    upload_id = uuid.uuid4().hex
    upload_dir = root_path / "data" / "user_groundtruth" / upload_id
    upload_path = upload_dir / filename
    upload_dir.mkdir(parents=True, exist_ok=False)
    try:
        upload_path.write_bytes(content)
        source_id = f"groundtruth:upload:{upload_id}"
        source = _groundtruth_record(
            source_id,
            label.strip() or Path(filename).stem,
            "uploaded",
            upload_path.resolve(),
        )
        if not source.valid:
            if source.validation == "missing_query_column":
                raise ValueError("Ground-truth upload is missing a supported query column")
            if source.validation == "missing_relevance_column":
                raise ValueError("Ground-truth upload is missing a supported relevance column")
            raise ValueError(f"Ground-truth upload is invalid: {source.validation}")
        manifest = {
            "label": source.label,
            "filename": filename,
            "row_count": source.row_count,
            "query_column": source.query_column,
            "answer_column": source.answer_column,
            "source_column": source.source_column,
        }
        (upload_dir / "manifest.json").write_text(
            json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
        payload = source.public()
        payload["source_id"] = source.id
        return payload
    except Exception:
        shutil.rmtree(upload_dir, ignore_errors=True)
        raise


def build_source_catalog(root: str | Path) -> dict[str, list[dict[str, Any]]]:
    datasets, groundtruth = _discover_sources(Path(root))
    return {
        "datasets": [source.public() for source in datasets],
        "groundtruth": [source.public() for source in groundtruth],
    }


def resolve_dataset(root: str | Path, source_id: str) -> DatasetSource:
    datasets, _groundtruth = _discover_sources(Path(root))
    source = next((item for item in datasets if item.id == source_id), None)
    if source is None:
        raise ValueError(f"Unknown source ID: {source_id}")
    return source


def resolve_groundtruth(root: str | Path, source_id: str) -> GroundtruthSource | None:
    if source_id == NONE_GROUNDTRUTH_ID:
        return None
    _datasets, groundtruth = _discover_sources(Path(root))
    source = next((item for item in groundtruth if item.id == source_id), None)
    if source is None or source.path is None:
        raise ValueError(f"Unknown source ID: {source_id}")
    if not source.valid:
        raise ValueError(f"Unknown or invalid source ID: {source_id}")
    return source
