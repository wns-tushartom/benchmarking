"""Read-only structural corpus admission; compose with approved hash checks."""
import csv
from pathlib import Path
from openpyxl import load_workbook
from benchmarking.core.config import technique
from benchmarking.adapters.local import load_query_cases, sniff_delimiter


def _path(root, value):
    relative = Path(value)
    path = root / relative
    if relative.is_absolute() or '..' in relative.parts or path.resolve() != path or not path.is_file():
        raise ValueError('corpus path rejected')
    return path


def verify_corpus(root, config, combinations):
    root = Path(root).absolute()
    experiment = config['experiment']
    dataset = _path(root, experiment['dataset'])
    workbook = _path(root, experiment['corpus_workbook'])
    with dataset.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(stream, delimiter=sniff_delimiter(dataset))
        headers = reader.fieldnames or []
        if len(headers) != len(set(headers)):
            raise ValueError('duplicate query columns')
        raw = list(reader)
    cases = load_query_cases(root, experiment['dataset'])
    if len(raw) != 500 or len(cases) != 500:
        raise ValueError('exactly 500 usable queries required')
    sheets = set()
    for row in combinations:
        params = technique(config, 'chunkers', row['chunker'])
        if params.get('adapter') != 'local_workbook':
            raise ValueError('unsupported corpus adapter')
        if _path(root, params.get('workbook', 'data/chunking_methods_output_v2.xlsx')) != workbook:
            raise ValueError('chunker workbook differs from bound corpus')
        sheets.add(params.get('sheet_name', row['chunker']))
    if not sheets:
        raise ValueError('no selected corpus sheets')
    evidence = []
    reference = None
    wb = load_workbook(workbook, read_only=True, data_only=True)
    try:
        for name in sorted(sheets):
            if name not in wb.sheetnames:
                raise ValueError('selected corpus sheet missing')
            rows = wb[name].iter_rows(values_only=True)
            headers = list(next(rows, ()))
            if len(headers) != len(set(headers)) or not {'id', 'pdf_name', 'paragraph'} <= set(headers):
                raise ValueError('invalid chunk schema')
            ids, pdfs = set(), set()
            for values in rows:
                row = dict(zip(headers, values))
                identity = row['id']
                if isinstance(identity, bool) or identity is None:
                    raise ValueError('invalid chunk identity')
                try:
                    numeric = int(identity)
                    if str(identity) != str(numeric) and identity != numeric:
                        raise ValueError('nonintegral chunk identity')
                except (TypeError, ValueError, OverflowError):
                    raise ValueError('invalid chunk identity') from None
                if numeric in ids:
                    raise ValueError('duplicate chunk identity')
                if not all(isinstance(row[k], str) and row[k].strip() for k in ('pdf_name', 'paragraph')):
                    raise ValueError('empty or invalid chunk content')
                ids.add(numeric); pdfs.add(row['pdf_name'])
            if not ids or (reference is not None and pdfs != reference):
                raise ValueError('empty or inconsistent corpus PDF coverage')
            reference = pdfs
            evidence.append({'sheet': name, 'row_count': len(ids), 'pdf_count': len(pdfs)})
    finally:
        wb.close()
    return {'corpus_verified': True, 'query_count': len(cases), 'pdf_count': len(reference), 'sheets': evidence}
