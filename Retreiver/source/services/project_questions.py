from __future__ import annotations

import csv
import ctypes
from dataclasses import dataclass, field
import errno
import hashlib
from io import BytesIO, StringIO
import json
import os
from pathlib import Path
import re
import shutil
import stat
import unicodedata
from typing import Any, Iterable
from uuid import uuid4

from source.services.project_chunking import OFFICIAL_PROJECT_CHUNKERS


_SCHEMA_VERSION = 1
_MAX_QUERY_CHARS = 4_000
_DEFAULT_MAX_QUESTIONS = 10_000
_SOURCE_ID_RE = re.compile(r"^source_[0-9a-f]{64}$")
_CHUNK_ID_RE = re.compile(r"^[A-Za-z0-9._-]+$")
_QUESTION_SET_FILE_SUFFIXES = {".txt", ".csv", ".xlsx"}
_TEXT_FORBIDDEN_SIGNATURES = (
    b"%PDF-",
    b"PK\x03\x04",
    b"PK\x05\x06",
    b"PK\x07\x08",
    b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",
    b"\x89PNG\r\n\x1a\n",
    b"GIF87a",
    b"GIF89a",
    b"\xff\xd8\xff",
)
_XLSX_SIGNATURE = b"PK\x03\x04"
_REQUIRED_ROW_KEYS = {
    "schema_version",
    "question_id",
    "query",
    "labels",
    "reference_answer",
    "source_row",
}
_REQUIRED_LABEL_KEYS = {"reference_contexts", "source_ids", "chunk_refs"}


class ProjectQuestionValidationError(ValueError):
    """Raised when a project question or immutable question set is invalid."""


class ProjectQuestionStorageError(RuntimeError):
    """Raised when a question set cannot be stored or loaded safely."""


@dataclass(frozen=True)
class QuestionLimits:
    max_questions: int = _DEFAULT_MAX_QUESTIONS
    max_file_bytes: int = 10 * 1024 * 1024

    def __post_init__(self) -> None:
        if (
            isinstance(self.max_questions, bool)
            or not isinstance(self.max_questions, int)
            or self.max_questions < 1
            or isinstance(self.max_file_bytes, bool)
            or not isinstance(self.max_file_bytes, int)
            or self.max_file_bytes < 1
        ):
            raise ProjectQuestionValidationError("question limit configuration is invalid")

    @classmethod
    def from_environment(cls) -> "QuestionLimits":
        try:
            max_questions = int(
                os.getenv("PROJECT_MAX_QUESTIONS", str(_DEFAULT_MAX_QUESTIONS))
            )
            max_file_kb = int(os.getenv("PROJECT_MAX_QUESTION_FILE_KB", "10240"))
        except (TypeError, ValueError, OverflowError):
            raise ProjectQuestionValidationError("question limit configuration is invalid") from None
        return cls(
            max_questions=max_questions,
            max_file_bytes=max_file_kb * 1024,
        )


def _sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _normalized_query_identity(query: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", query).casefold().split())


def _validate_query(value: Any) -> str:
    if not isinstance(value, str):
        raise ProjectQuestionValidationError("question must be text")
    query = value.strip()
    if not query:
        raise ProjectQuestionValidationError("question must not be empty")
    if len(query) > _MAX_QUERY_CHARS:
        raise ProjectQuestionValidationError("question exceeds the 4,000 character limit")
    return query


@dataclass(frozen=True)
class ChunkRef:
    chunker: str
    chunk_id: str

    def __post_init__(self) -> None:
        if self.chunker not in OFFICIAL_PROJECT_CHUNKERS:
            raise ProjectQuestionValidationError("chunk_ref uses an unknown chunker")
        if not isinstance(self.chunk_id, str) or _CHUNK_ID_RE.fullmatch(self.chunk_id) is None:
            raise ProjectQuestionValidationError("chunk_ref has an invalid chunk id")

    def to_row(self) -> dict[str, str]:
        return {"chunker": self.chunker, "chunk_id": self.chunk_id}


@dataclass(frozen=True)
class RetrievalLabels:
    reference_contexts: tuple[str, ...] = ()
    source_ids: tuple[str, ...] = ()
    chunk_refs: tuple[ChunkRef, ...] = ()

    def __post_init__(self) -> None:
        contexts = tuple(self.reference_contexts)
        source_ids = tuple(self.source_ids)
        chunk_refs = tuple(self.chunk_refs)
        if any(not isinstance(value, str) or not value.strip() for value in contexts):
            raise ProjectQuestionValidationError("reference context must be non-empty text")
        if any(_SOURCE_ID_RE.fullmatch(value) is None for value in source_ids):
            raise ProjectQuestionValidationError("source_id is invalid")
        if any(not isinstance(value, ChunkRef) for value in chunk_refs):
            raise ProjectQuestionValidationError("chunk_ref is invalid")
        object.__setattr__(self, "reference_contexts", contexts)
        object.__setattr__(self, "source_ids", source_ids)
        object.__setattr__(self, "chunk_refs", chunk_refs)

    def is_empty(self) -> bool:
        return not (self.reference_contexts or self.source_ids or self.chunk_refs)

    def to_row(self) -> dict[str, Any]:
        return {
            "reference_contexts": list(self.reference_contexts),
            "source_ids": list(self.source_ids),
            "chunk_refs": [reference.to_row() for reference in self.chunk_refs],
        }


@dataclass(frozen=True)
class ProjectQuestion:
    question_id: str
    query: str
    labels: RetrievalLabels = field(default_factory=RetrievalLabels)
    reference_answer: str | None = None
    source_row: int = 1

    def __post_init__(self) -> None:
        if re.fullmatch(r"q_[0-9]{6}", self.question_id or "") is None:
            raise ProjectQuestionValidationError("question_id is invalid")
        query = _validate_query(self.query)
        if not isinstance(self.labels, RetrievalLabels):
            raise ProjectQuestionValidationError("labels are invalid")
        answer = self.reference_answer
        if answer is not None:
            if not isinstance(answer, str):
                raise ProjectQuestionValidationError("reference answer must be text or null")
            answer = answer.strip() or None
        if isinstance(self.source_row, bool) or not isinstance(self.source_row, int) or self.source_row < 1:
            raise ProjectQuestionValidationError("source_row is invalid")
        object.__setattr__(self, "query", query)
        object.__setattr__(self, "reference_answer", answer)

    def to_row(self) -> dict[str, Any]:
        return {
            "schema_version": _SCHEMA_VERSION,
            "question_id": self.question_id,
            "query": self.query,
            "labels": self.labels.to_row(),
            "reference_answer": self.reference_answer,
            "source_row": self.source_row,
        }


def _dedupe(values: Iterable[str]) -> tuple[str, ...]:
    output: list[str] = []
    seen: set[str] = set()
    for value in values:
        stripped = value.strip()
        if stripped and stripped not in seen:
            seen.add(stripped)
            output.append(stripped)
    return tuple(output)


def _list_cell(value: Any) -> tuple[str, ...]:
    if value is None:
        return ()
    if not isinstance(value, str):
        value = str(value)
    stripped = value.strip()
    if not stripped:
        return ()
    if stripped.startswith("["):
        try:
            decoded = json.loads(stripped)
        except json.JSONDecodeError:
            raise ProjectQuestionValidationError("list-valued label is malformed") from None
        if not isinstance(decoded, list) or any(not isinstance(item, str) for item in decoded):
            raise ProjectQuestionValidationError("list-valued label must be a text array")
        return _dedupe(decoded)
    return _dedupe(stripped.split(";"))


def _parse_source_ids(value: Any) -> tuple[str, ...]:
    values = _list_cell(value)
    if any(_SOURCE_ID_RE.fullmatch(item) is None for item in values):
        raise ProjectQuestionValidationError("source_id is invalid")
    return values


def _parse_chunk_refs(value: Any) -> tuple[ChunkRef, ...]:
    references: list[ChunkRef] = []
    seen: set[tuple[str, str]] = set()
    for item in _list_cell(value):
        if ":" not in item:
            raise ProjectQuestionValidationError(
                "chunk_ref must be qualified as <canonical_chunker_id>:<chunk_id>"
            )
        chunker, chunk_id = item.split(":", 1)
        try:
            reference = ChunkRef(chunker=chunker.strip(), chunk_id=chunk_id.strip())
        except ProjectQuestionValidationError:
            raise ProjectQuestionValidationError(
                "chunk_ref must be qualified as <canonical_chunker_id>:<chunk_id>"
            ) from None
        identity = (reference.chunker, reference.chunk_id)
        if identity not in seen:
            seen.add(identity)
            references.append(reference)
    return tuple(references)


def _header_map(values: list[Any]) -> dict[str, int]:
    headers: dict[str, int] = {}
    for index, value in enumerate(values):
        header = str(value or "").strip().casefold()
        if not header:
            continue
        if header in headers:
            raise ProjectQuestionValidationError(f"duplicate question column: {header}")
        headers[header] = index
    query_columns = [column for column in ("question", "query") if column in headers]
    if len(query_columns) != 1:
        raise ProjectQuestionValidationError("exactly one question or query column is required")
    for aliases in (("reference_context", "context"), ("answer", "ground_truth")):
        if sum(alias in headers for alias in aliases) > 1:
            raise ProjectQuestionValidationError(f"ambiguous columns: {aliases[0]} and {aliases[1]}")
    return headers


def _value(row: list[Any], headers: dict[str, int], *names: str) -> Any:
    for name in names:
        index = headers.get(name)
        if index is not None:
            return row[index] if index < len(row) else None
    return None


def _questions_from_rows(
    rows: Iterable[tuple[int, list[Any]]],
    *,
    max_questions: int,
) -> list[ProjectQuestion]:
    row_iterator = iter(rows)
    try:
        _, header_values = next(row_iterator)
    except StopIteration:
        raise ProjectQuestionValidationError("question file is empty") from None
    headers = _header_map(header_values)
    query_column = "question" if "question" in headers else "query"
    questions: list[ProjectQuestion] = []
    seen_queries: set[str] = set()
    max_physical_rows = max_questions + 1  # one header plus bounded data rows
    for source_row, row in row_iterator:
        if source_row > max_physical_rows:
            raise ProjectQuestionValidationError(
                "question file exceeds the maximum physical row count"
            )
        if not any(value is not None and str(value).strip() for value in row):
            continue
        query = _validate_query(_value(row, headers, query_column))
        identity = _normalized_query_identity(query)
        if identity in seen_queries:
            raise ProjectQuestionValidationError(f"duplicate question at source row {source_row}")
        seen_queries.add(identity)
        contexts = _list_cell(_value(row, headers, "reference_context", "context"))
        source_ids = _parse_source_ids(_value(row, headers, "source_id"))
        chunk_refs = _parse_chunk_refs(_value(row, headers, "chunk_ref"))
        answer_value = _value(row, headers, "answer", "ground_truth")
        reference_answer = str(answer_value).strip() if answer_value is not None else None
        questions.append(
            ProjectQuestion(
                question_id=f"q_{len(questions) + 1:06d}",
                query=query,
                labels=RetrievalLabels(
                    reference_contexts=contexts,
                    source_ids=source_ids,
                    chunk_refs=chunk_refs,
                ),
                reference_answer=reference_answer or None,
                source_row=source_row,
            )
        )
        if len(questions) > max_questions:
            raise ProjectQuestionValidationError("question file exceeds the maximum question count")
    if not questions:
        raise ProjectQuestionValidationError("question file contains no questions")
    return questions


def parse_typed_question(query: str) -> list[ProjectQuestion]:
    from source.services.project_relevance import load_pinned_stopwords

    load_pinned_stopwords()
    return [
        ProjectQuestion(
            question_id="q_000001",
            query=query,
            labels=RetrievalLabels(),
            reference_answer=None,
            source_row=1,
        )
    ]


def _parse_txt(content: bytes, max_questions: int) -> list[ProjectQuestion]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeError:
        raise ProjectQuestionValidationError("TXT question file must be UTF-8") from None
    questions: list[ProjectQuestion] = []
    seen: set[str] = set()
    for source_row, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            continue
        query = _validate_query(line)
        identity = _normalized_query_identity(query)
        if identity in seen:
            raise ProjectQuestionValidationError(f"duplicate question at source row {source_row}")
        seen.add(identity)
        questions.append(
            ProjectQuestion(
                question_id=f"q_{len(questions) + 1:06d}",
                query=query,
                source_row=source_row,
            )
        )
        if len(questions) > max_questions:
            raise ProjectQuestionValidationError("question file exceeds the maximum question count")
    if not questions:
        raise ProjectQuestionValidationError("question file contains no questions")
    return questions


def _parse_csv(content: bytes, max_questions: int) -> list[ProjectQuestion]:
    try:
        text = content.decode("utf-8-sig")
        rows = (
            (source_row, list(row))
            for source_row, row in enumerate(
                csv.reader(StringIO(text, newline=""), strict=True), 1
            )
        )
        return _questions_from_rows(rows, max_questions=max_questions)
    except ProjectQuestionValidationError:
        raise
    except (UnicodeError, csv.Error):
        raise ProjectQuestionValidationError("CSV question file is malformed") from None


def _parse_xlsx(content: bytes, max_questions: int) -> list[ProjectQuestion]:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if sheet is None:
                raise ProjectQuestionValidationError("XLSX question file has no active sheet")
            if sheet.max_row > max_questions + 1:
                raise ProjectQuestionValidationError(
                    "question file exceeds the maximum physical row count"
                )
            rows = (
                (source_row, list(row))
                for source_row, row in enumerate(sheet.iter_rows(values_only=True), 1)
            )
            return _questions_from_rows(rows, max_questions=max_questions)
        finally:
            workbook.close()
    except ProjectQuestionValidationError:
        raise
    except Exception:
        raise ProjectQuestionValidationError("XLSX question file is malformed") from None


def parse_question_bytes(
    filename: str,
    content: bytes,
    *,
    max_questions: int | None = None,
    max_file_bytes: int | None = None,
    upload_limits: Any | None = None,
) -> list[ProjectQuestion]:
    from source.services.project_relevance import load_pinned_stopwords
    from source.services.project_workspace import (
        UploadError,
        UploadLimits,
        _validate_payload_signature,
    )

    load_pinned_stopwords()
    environment_limits = QuestionLimits.from_environment()
    if max_questions is not None and (
        isinstance(max_questions, bool)
        or not isinstance(max_questions, int)
        or max_questions < 1
    ):
        raise ProjectQuestionValidationError("question limit configuration is invalid")
    if max_file_bytes is not None and (
        isinstance(max_file_bytes, bool)
        or not isinstance(max_file_bytes, int)
        or max_file_bytes < 1
    ):
        raise ProjectQuestionValidationError("question limit configuration is invalid")
    effective_limits = QuestionLimits(
        max_questions=(
            environment_limits.max_questions
            if max_questions is None
            else min(environment_limits.max_questions, max_questions)
        ),
        max_file_bytes=(
            environment_limits.max_file_bytes
            if max_file_bytes is None
            else min(environment_limits.max_file_bytes, max_file_bytes)
        ),
    )
    if not isinstance(filename, str) or Path(filename).name != filename or not filename.strip():
        raise ProjectQuestionValidationError("question filename is invalid")
    if not isinstance(content, bytes) or not content:
        raise ProjectQuestionValidationError("question file is empty")
    if len(content) > effective_limits.max_file_bytes:
        raise ProjectQuestionValidationError("question file exceeds the configured size limit")
    suffix = Path(filename).suffix.lower()
    if suffix not in _QUESTION_SET_FILE_SUFFIXES:
        raise ProjectQuestionValidationError("question file type is unsupported")
    effective_upload_limits = UploadLimits() if upload_limits is None else upload_limits
    if not isinstance(effective_upload_limits, UploadLimits):
        raise ProjectQuestionValidationError("question limit configuration is invalid")
    try:
        _validate_payload_signature(filename, content, effective_upload_limits)
    except UploadError:
        raise ProjectQuestionValidationError(
            "question file signature or container is invalid"
        ) from None
    if suffix in {".txt", ".csv"}:
        if b"\x00" in content or any(
            content.startswith(signature) for signature in _TEXT_FORBIDDEN_SIGNATURES
        ):
            raise ProjectQuestionValidationError("question file signature does not match its extension")
    elif not content.startswith(_XLSX_SIGNATURE):
        raise ProjectQuestionValidationError("question file signature does not match its extension")
    if suffix == ".txt":
        return _parse_txt(content, effective_limits.max_questions)
    if suffix == ".csv":
        return _parse_csv(content, effective_limits.max_questions)
    return _parse_xlsx(content, effective_limits.max_questions)


def _render_questions(questions: Iterable[ProjectQuestion]) -> bytes:
    rows = list(questions)
    if not rows:
        raise ProjectQuestionValidationError("question set is empty")
    return b"".join(
        json.dumps(question.to_row(), sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        + b"\n"
        for question in rows
    )


def _labels_from_row(value: Any) -> RetrievalLabels:
    if not isinstance(value, dict) or set(value) != _REQUIRED_LABEL_KEYS:
        raise ProjectQuestionValidationError("normalized question label schema is invalid")
    contexts = value["reference_contexts"]
    source_ids = value["source_ids"]
    chunk_refs = value["chunk_refs"]
    if not isinstance(contexts, list) or not isinstance(source_ids, list) or not isinstance(chunk_refs, list):
        raise ProjectQuestionValidationError("normalized question label schema is invalid")
    if any(not isinstance(item, str) for item in contexts + source_ids):
        raise ProjectQuestionValidationError("normalized question labels are invalid")
    if tuple(contexts) != _dedupe(contexts) or tuple(source_ids) != _dedupe(source_ids):
        raise ProjectQuestionValidationError("normalized question labels are not canonical")
    references: list[ChunkRef] = []
    seen_references: set[tuple[str, str]] = set()
    for item in chunk_refs:
        if not isinstance(item, dict) or set(item) != {"chunker", "chunk_id"}:
            raise ProjectQuestionValidationError("normalized chunk_ref schema is invalid")
        reference = ChunkRef(chunker=item["chunker"], chunk_id=item["chunk_id"])
        identity = (reference.chunker, reference.chunk_id)
        if identity in seen_references:
            raise ProjectQuestionValidationError("normalized question labels are not canonical")
        seen_references.add(identity)
        references.append(reference)
    labels = RetrievalLabels(tuple(contexts), tuple(source_ids), tuple(references))
    if labels.to_row() != value:
        raise ProjectQuestionValidationError("normalized question labels are not canonical")
    return labels


def _read_regular_bytes(path: Path, description: str) -> bytes:
    path = Path(path)
    try:
        for component in path.parents:
            metadata = component.lstat()
            if stat.S_ISLNK(metadata.st_mode):
                raise ProjectQuestionStorageError(f"{description} is unavailable")
        parent = path.parent
        flags = os.O_RDONLY | getattr(os, "O_DIRECTORY", 0) | getattr(os, "O_NOFOLLOW", 0)
        parent_fd = os.open(parent, flags)
        try:
            opened_parent = os.fstat(parent_fd)
            lexical_parent = parent.lstat()
            if (
                not stat.S_ISDIR(opened_parent.st_mode)
                or not stat.S_ISDIR(lexical_parent.st_mode)
                or (opened_parent.st_dev, opened_parent.st_ino)
                != (lexical_parent.st_dev, lexical_parent.st_ino)
                or parent.resolve(strict=True) != parent
            ):
                raise ProjectQuestionStorageError(f"{description} is unavailable")
            content = _read_regular_bytes_at(parent_fd, path.name, description)
            current_parent = parent.lstat()
            if (
                stat.S_ISLNK(current_parent.st_mode)
                or (current_parent.st_dev, current_parent.st_ino)
                != (opened_parent.st_dev, opened_parent.st_ino)
            ):
                raise ProjectQuestionStorageError(f"{description} changed during loading")
            return content
        finally:
            os.close(parent_fd)
    except ProjectQuestionStorageError:
        raise
    except OSError:
        raise ProjectQuestionStorageError(f"{description} is unavailable") from None


def _read_regular_bytes_at(parent_fd: int, name: str, description: str) -> bytes:
    if not isinstance(name, str) or Path(name).name != name:
        raise ProjectQuestionStorageError(f"{description} is unavailable")
    flags = os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0)
    try:
        metadata = os.stat(name, dir_fd=parent_fd, follow_symlinks=False)
        if stat.S_ISLNK(metadata.st_mode) or not stat.S_ISREG(metadata.st_mode):
            raise ProjectQuestionStorageError(f"{description} is unavailable")
        descriptor = os.open(name, flags, dir_fd=parent_fd)
    except ProjectQuestionStorageError:
        raise
    except OSError:
        raise ProjectQuestionStorageError(f"{description} is unavailable") from None
    try:
        opened = os.fstat(descriptor)
        if (
            not stat.S_ISREG(opened.st_mode)
            or (opened.st_dev, opened.st_ino) != (metadata.st_dev, metadata.st_ino)
        ):
            raise ProjectQuestionStorageError(f"{description} is unavailable")
        chunks: list[bytes] = []
        while True:
            chunk = os.read(descriptor, 1024 * 1024)
            if not chunk:
                break
            chunks.append(chunk)
        return b"".join(chunks)
    finally:
        os.close(descriptor)


def _parse_normalized_questions(
    content: bytes,
    *,
    expected_sha256: str | None = None,
) -> list[ProjectQuestion]:
    if expected_sha256 is not None and _sha256(content) != expected_sha256:
        raise ProjectQuestionValidationError("normalized question hash mismatch")
    questions: list[ProjectQuestion] = []
    seen_ids: set[str] = set()
    seen_queries: set[str] = set()
    try:
        for line in content.decode("utf-8").splitlines():
            row = json.loads(line)
            if not isinstance(row, dict) or set(row) != _REQUIRED_ROW_KEYS:
                raise ProjectQuestionValidationError("normalized question schema is invalid")
            if type(row["schema_version"]) is not int or row["schema_version"] != _SCHEMA_VERSION:
                raise ProjectQuestionValidationError("normalized question schema is invalid")
            expected_question_id = f"q_{len(questions) + 1:06d}"
            if row["question_id"] != expected_question_id:
                raise ProjectQuestionValidationError("normalized questions are not canonical")
            question = ProjectQuestion(
                question_id=row["question_id"],
                query=row["query"],
                labels=_labels_from_row(row["labels"]),
                reference_answer=row["reference_answer"],
                source_row=row["source_row"],
            )
            if question.question_id in seen_ids or _normalized_query_identity(question.query) in seen_queries:
                raise ProjectQuestionValidationError("normalized question contains duplicates")
            seen_ids.add(question.question_id)
            seen_queries.add(_normalized_query_identity(question.query))
            questions.append(question)
    except (UnicodeError, json.JSONDecodeError, TypeError, KeyError):
        raise ProjectQuestionValidationError("normalized question schema is invalid") from None
    if not questions:
        raise ProjectQuestionValidationError("question set is empty")
    if _render_questions(questions) != content:
        raise ProjectQuestionValidationError("normalized questions are not canonical")
    return questions


def load_normalized_questions(
    path: Path,
    *,
    expected_sha256: str | None = None,
) -> list[ProjectQuestion]:
    from source.services.project_relevance import load_pinned_stopwords

    load_pinned_stopwords()
    path = Path(path)
    if path.name != "questions.jsonl":
        raise ProjectQuestionValidationError("canonical questions.jsonl is required")
    return _parse_normalized_questions(
        _read_regular_bytes(path, "canonical questions.jsonl"),
        expected_sha256=expected_sha256,
    )


def _fsync_directory(path: Path) -> None:
    descriptor = os.open(path, os.O_RDONLY | getattr(os, "O_DIRECTORY", 0))
    try:
        os.fsync(descriptor)
    finally:
        os.close(descriptor)


def _write_new(path: Path, content: bytes) -> None:
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(descriptor, "wb") as output:
        output.write(content)
        output.flush()
        os.fsync(output.fileno())


def _validate_question_set_id(value: str) -> str:
    from source.services.project_workspace import validate_resource_id

    try:
        validated = validate_resource_id(value)
    except (TypeError, ValueError):
        raise ProjectQuestionValidationError("question_set_id is invalid") from None
    if not validated.startswith("questions_"):
        raise ProjectQuestionValidationError("question_set_id is invalid")
    return validated


def _publish_directory_noreplace(parent: Path, staging_name: str, final_name: str) -> None:
    flags = os.O_RDONLY | getattr(os, "O_DIRECTORY", 0) | getattr(os, "O_NOFOLLOW", 0)
    try:
        parent_fd = os.open(parent, flags)
    except OSError:
        raise ProjectQuestionStorageError("question storage is unavailable") from None
    try:
        opened_parent = os.fstat(parent_fd)
        lexical_parent = os.stat(parent, follow_symlinks=False)
        if (
            not stat.S_ISDIR(opened_parent.st_mode)
            or not stat.S_ISDIR(lexical_parent.st_mode)
            or (opened_parent.st_dev, opened_parent.st_ino)
            != (lexical_parent.st_dev, lexical_parent.st_ino)
            or parent.resolve(strict=True) != parent
        ):
            raise ProjectQuestionStorageError("question storage is unavailable")
        try:
            staging_metadata = os.stat(
                staging_name,
                dir_fd=parent_fd,
                follow_symlinks=False,
            )
        except OSError:
            raise ProjectQuestionStorageError("question set could not be stored") from None
        if not stat.S_ISDIR(staging_metadata.st_mode):
            raise ProjectQuestionStorageError("question set could not be stored")

        libc = ctypes.CDLL(None, use_errno=True)
        renameat2 = getattr(libc, "renameat2", None)
        if renameat2 is None:
            raise ProjectQuestionStorageError("question set could not be stored")
        renameat2.argtypes = [ctypes.c_int, ctypes.c_char_p, ctypes.c_int, ctypes.c_char_p, ctypes.c_uint]
        renameat2.restype = ctypes.c_int
        result = renameat2(
            parent_fd,
            os.fsencode(staging_name),
            parent_fd,
            os.fsencode(final_name),
            1,  # RENAME_NOREPLACE
        )
        if result != 0:
            error_number = ctypes.get_errno()
            if error_number == errno.EEXIST:
                raise ProjectQuestionStorageError("question set already exists")
            raise ProjectQuestionStorageError("question set could not be stored")
        try:
            os.fsync(parent_fd)
        except OSError:
            try:
                os.rename(
                    final_name,
                    staging_name,
                    src_dir_fd=parent_fd,
                    dst_dir_fd=parent_fd,
                )
            except OSError:
                raise ProjectQuestionStorageError(
                    "question set publication state is indeterminate"
                ) from None
            try:
                os.fsync(parent_fd)
            except OSError:
                raise ProjectQuestionStorageError(
                    "question set publication state is indeterminate"
                ) from None
            raise ProjectQuestionStorageError("question set could not be stored") from None
    finally:
        os.close(parent_fd)


def create_question_set(
    questions_root: Path,
    question_set_id: str,
    filename: str,
    content: bytes,
    *,
    max_questions: int | None = None,
    upload_limits: Any | None = None,
) -> dict[str, Any]:
    question_set_id = _validate_question_set_id(question_set_id)
    questions = parse_question_bytes(
        filename,
        content,
        max_questions=max_questions,
        upload_limits=upload_limits,
    )
    normalized = _render_questions(questions)
    content_sha256 = _sha256(content)
    normalized_sha256 = _sha256(normalized)
    questions_root = Path(questions_root)
    if questions_root.is_symlink() or not questions_root.is_dir() or questions_root.resolve() != questions_root:
        raise ProjectQuestionStorageError("question storage is unavailable")
    final_root = questions_root / question_set_id
    staging = questions_root / f".staging-{uuid4().hex}"
    if final_root.exists() or final_root.is_symlink():
        raise ProjectQuestionStorageError("question set already exists")
    created = False
    try:
        staging.mkdir(mode=0o700)
        created = True
        _write_new(staging / filename, content)
        _write_new(staging / "questions.jsonl", normalized)
        manifest = {
            "schema_version": _SCHEMA_VERSION,
            "question_set_id": question_set_id,
            "source_filename": filename,
            "content_sha256": content_sha256,
            "normalized_sha256": normalized_sha256,
            "question_count": len(questions),
        }
        _write_new(
            staging / "manifest.json",
            (json.dumps(manifest, sort_keys=True, separators=(",", ":")) + "\n").encode("utf-8"),
        )
        _fsync_directory(staging)
        _publish_directory_noreplace(
            questions_root,
            staging.name,
            final_root.name,
        )
        created = False
        return manifest
    except (ProjectQuestionValidationError, ProjectQuestionStorageError):
        raise
    except OSError:
        raise ProjectQuestionStorageError("question set could not be stored") from None
    finally:
        if created and staging.parent == questions_root and staging.exists() and not staging.is_symlink():
            shutil.rmtree(staging)


def _open_question_set_directory(
    questions_root: Path,
    question_set_id: str,
) -> tuple[int, int]:
    flags = os.O_RDONLY | getattr(os, "O_DIRECTORY", 0) | getattr(os, "O_NOFOLLOW", 0)
    parent_fd = -1
    try:
        parent_metadata = questions_root.lstat()
        if stat.S_ISLNK(parent_metadata.st_mode) or not stat.S_ISDIR(parent_metadata.st_mode):
            raise ProjectQuestionStorageError("question set is unavailable")
        parent_fd = os.open(questions_root, flags)
        opened_parent = os.fstat(parent_fd)
        if (
            not stat.S_ISDIR(opened_parent.st_mode)
            or (opened_parent.st_dev, opened_parent.st_ino)
            != (parent_metadata.st_dev, parent_metadata.st_ino)
            or questions_root.resolve(strict=True) != questions_root
        ):
            raise ProjectQuestionStorageError("question set is unavailable")
        root_metadata = os.stat(
            question_set_id,
            dir_fd=parent_fd,
            follow_symlinks=False,
        )
        if stat.S_ISLNK(root_metadata.st_mode) or not stat.S_ISDIR(root_metadata.st_mode):
            raise ProjectQuestionStorageError("question set is unavailable")
        root_fd = os.open(question_set_id, flags, dir_fd=parent_fd)
        opened_root = os.fstat(root_fd)
        if (
            not stat.S_ISDIR(opened_root.st_mode)
            or (opened_root.st_dev, opened_root.st_ino)
            != (root_metadata.st_dev, root_metadata.st_ino)
        ):
            os.close(root_fd)
            raise ProjectQuestionStorageError("question set is unavailable")
        return parent_fd, root_fd
    except ProjectQuestionStorageError:
        if parent_fd >= 0:
            os.close(parent_fd)
        raise
    except FileNotFoundError:
        if parent_fd >= 0:
            os.close(parent_fd)
        raise FileNotFoundError("question set does not exist") from None
    except OSError:
        if parent_fd >= 0:
            os.close(parent_fd)
        raise ProjectQuestionStorageError("question set is unavailable") from None


def _verify_open_directory_unchanged(
    questions_root: Path,
    question_set_id: str,
    parent_fd: int,
    root_fd: int,
) -> None:
    try:
        current_parent = questions_root.lstat()
        opened_parent = os.fstat(parent_fd)
        current_root = os.stat(
            question_set_id,
            dir_fd=parent_fd,
            follow_symlinks=False,
        )
        opened_root = os.fstat(root_fd)
        if (
            stat.S_ISLNK(current_parent.st_mode)
            or not stat.S_ISDIR(current_parent.st_mode)
            or (current_parent.st_dev, current_parent.st_ino)
            != (opened_parent.st_dev, opened_parent.st_ino)
            or stat.S_ISLNK(current_root.st_mode)
            or not stat.S_ISDIR(current_root.st_mode)
            or (current_root.st_dev, current_root.st_ino)
            != (opened_root.st_dev, opened_root.st_ino)
        ):
            raise ProjectQuestionStorageError("question set changed during loading")
    except ProjectQuestionStorageError:
        raise
    except OSError:
        raise ProjectQuestionStorageError("question set changed during loading") from None


def load_question_set(
    questions_root: Path,
    question_set_id: str,
    *,
    expected_content_sha256: str,
) -> list[ProjectQuestion]:
    question_set_id = _validate_question_set_id(question_set_id)
    questions_root = Path(questions_root)
    root = questions_root / question_set_id
    if (
        questions_root.is_symlink()
        or not questions_root.is_dir()
        or questions_root.resolve() != questions_root
        or root.parent != questions_root
    ):
        raise FileNotFoundError("question set does not exist")
    parent_fd, root_fd = _open_question_set_directory(
        questions_root,
        question_set_id,
    )
    try:
        try:
            manifest = json.loads(
                _read_regular_bytes_at(
                    root_fd, "manifest.json", "question set manifest"
                ).decode("utf-8")
            )
        except ProjectQuestionStorageError:
            raise
        except (UnicodeError, json.JSONDecodeError):
            raise ProjectQuestionStorageError("question set manifest is unavailable") from None
        required = {
            "schema_version",
            "question_set_id",
            "source_filename",
            "content_sha256",
            "normalized_sha256",
            "question_count",
        }
        if not isinstance(manifest, dict) or set(manifest) != required:
            raise ProjectQuestionValidationError("question set manifest schema is invalid")
        if (
            type(manifest["schema_version"]) is not int
            or manifest["schema_version"] != _SCHEMA_VERSION
            or manifest["question_set_id"] != question_set_id
            or not isinstance(manifest["content_sha256"], str)
            or re.fullmatch(r"[0-9a-f]{64}", manifest["content_sha256"]) is None
            or not isinstance(manifest["normalized_sha256"], str)
            or re.fullmatch(r"[0-9a-f]{64}", manifest["normalized_sha256"]) is None
            or isinstance(manifest["question_count"], bool)
            or not isinstance(manifest["question_count"], int)
            or manifest["question_count"] < 1
        ):
            raise ProjectQuestionValidationError("question set manifest schema is invalid")
        filename = manifest["source_filename"]
        if not isinstance(filename, str) or Path(filename).name != filename:
            raise ProjectQuestionValidationError("question set manifest schema is invalid")
        if not isinstance(expected_content_sha256, str) or expected_content_sha256 != manifest["content_sha256"]:
            raise ProjectQuestionValidationError("question set content hash mismatch")
        raw_content = _read_regular_bytes_at(root_fd, filename, "question set source")
        if _sha256(raw_content) != expected_content_sha256:
            raise ProjectQuestionValidationError("question set content hash mismatch")
        normalized_content = _read_regular_bytes_at(
            root_fd, "questions.jsonl", "canonical questions.jsonl"
        )
        reparsed_content = _render_questions(
            parse_question_bytes(filename, raw_content)
        )
        if normalized_content != reparsed_content:
            raise ProjectQuestionValidationError(
                "normalized questions do not match the immutable source"
            )
        questions = _parse_normalized_questions(
            normalized_content,
            expected_sha256=manifest["normalized_sha256"],
        )
        if len(questions) != manifest["question_count"]:
            raise ProjectQuestionValidationError("question set count mismatch")
        _verify_open_directory_unchanged(
            questions_root,
            question_set_id,
            parent_fd,
            root_fd,
        )
        return questions
    finally:
        os.close(root_fd)
        os.close(parent_fd)
